"""Étape É1 — matrice PLU de Sète, générée (jamais saisie à la main).

Priorité des sources (décision 2026-09-21) :
  1. XML du règlement : taux d'emprise, mode de règle, pages, confiance (fait foi) ;
  2. Excel de classement : lecture qualitative (habitation, profil, contraintes, priorité de zone) ;
  3. config/familles.yaml : famille d'opération.

Sorties : config/plu_sete.yaml, web/data/regles.json, config/rapport_regles.md.
Lancer : python -m pipeline.build_regles   (depuis VIGIE/)
"""
from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import yaml

RACINE = Path(__file__).resolve().parents[1]
XML = RACINE / "sources" / "reglement_plu_sete_par_zone.xml"
XLSX = RACINE / "sources" / "classement_prospection_fonciere_sete.xlsx"
FAMILLES = RACINE / "config" / "familles.yaml"

# Modes du XML pour lesquels aucun taux applicable n'existe : réserve « non calculable ».
MODES_NON_CHIFFRES = {"formule_spatiale", "non_reglementee", "document_graphique", "non_explicite"}


def lire_xml() -> tuple[dict, dict]:
    """Retourne (secteurs, meta). secteurs[code] = {zone, categorie, emprise, piscine}."""
    racine = ET.parse(XML).getroot()
    secteurs: dict[str, dict] = {}
    for zone in racine.iter("zone"):
        for s in zone.find("sectors"):
            secteurs[s.get("code")] = {
                "zone": zone.get("code"),
                "categorie": zone.get("category"),
                "pages_zone": f"{zone.get('source_page_start')}-{zone.get('source_page_end')}",
                "emprise": None,
                "piscine": None,
            }
        for regle in zone.find("normalized_data").iter("rule"):
            mode = regle.get("mode")
            taux = float(regle.get("max_ratio")) if regle.get("max_ratio") else None
            info = {
                "mode": mode,
                "taux": taux,
                "texte": " ".join((regle.text or "").split()),
                "pages": regle.get("source_pages"),
                "confiance": regle.get("confidence"),
            }
            for code in regle.get("sectors", "").split():
                # Secteurs cités par une règle mais absents de la liste (ex. UEc1, UC5a).
                cible = secteurs.setdefault(code, {
                    "zone": zone.get("code"), "categorie": zone.get("category"),
                    "pages_zone": f"{zone.get('source_page_start')}-{zone.get('source_page_end')}",
                    "emprise": None, "piscine": None, "hors_preambule": True,
                })
                if mode == "majoration_piscine":
                    cible["piscine"] = info
                else:
                    cible["emprise"] = info
    meta = {
        "document": racine.findtext("metadata/document_title"),
        "fichier": racine.findtext("metadata/source_filename"),
        "sha256": racine.findtext("metadata/source_sha256"),
    }
    return secteurs, meta


def lire_xlsx() -> tuple[dict, list]:
    """Retourne (qualitatif par code, variantes) ; « UA5a alignement/recul » → variantes de UA5a."""
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Classement"]
    entetes = [c.value for c in ws[6]]
    col = {nom: i for i, nom in enumerate(entetes)}
    qualitatif: dict[str, dict] = {}
    for ligne in ws.iter_rows(min_row=7, values_only=True):
        libelle = ligne[col["Secteur"]]
        if not libelle:
            continue
        code, _, variante = str(libelle).partition(" ")
        entree = {
            "habitation": ligne[col["Habitation standard"]],
            "emprise_xlsx": ligne[col["Emprise maximale"]],
            "regle_emprise": ligne[col["Règle d'emprise"]],
            "profil": ligne[col["Profil d'opération"]],
            "forme_urbaine": ligne[col["Forme urbaine / usage"]],
            "hauteur": ligne[col["Hauteur repère"]],
            "contraintes": ligne[col["Contraintes principales"]],
            "score_zone": ligne[col["Score /100"]],
            "priorite_zone": ligne[col["Priorité"]],
            "controle": ligne[col["Contrôle préalable"]],
            "pages_excel": ligne[col["Pages source"]],
        }
        if variante:
            qualitatif.setdefault(code, {"variantes": {}})["variantes"][variante] = entree
        else:
            qualitatif[code] = entree
    # UA5a : pas de ligne unique → on retient la variante « alignement » comme référence qualitative.
    for code, q in list(qualitatif.items()):
        if "variantes" in q and len(q) == 1:
            ref = q["variantes"].get("alignement") or next(iter(q["variantes"].values()))
            qualitatif[code] = {**ref, "variantes": q["variantes"]}

    anomalies = []
    ws_m = wb["Méthode"]
    dans_bloc = False
    for ligne in ws_m.iter_rows(values_only=True):
        valeurs = [v for v in ligne if v is not None]
        if not valeurs:
            continue
        if str(valeurs[0]).startswith("Anomalies"):
            dans_bloc = True
            continue
        if dans_bloc:
            if len(valeurs) < 2:
                break
            anomalies.append({"secteurs": str(valeurs[0]), "note": str(valeurs[1])})
    return qualitatif, anomalies


def famille_de(code: str, familles: dict) -> str:
    if code.endswith("v"):
        return "protege"
    for cle, f in familles.items():
        if code in f["secteurs"]:
            return cle
    return "non_classe"


def construire() -> dict:
    secteurs, meta = lire_xml()
    qualitatif, anomalies = lire_xlsx()
    familles = yaml.safe_load(FAMILLES.read_text(encoding="utf-8"))["familles"]

    ecarts, sans_excel = [], []
    sortie: dict[str, dict] = {}
    for code in sorted(secteurs):
        s = secteurs[code]
        q = qualitatif.get(code)
        emprise = s["emprise"] or {"mode": "absente", "taux": None, "texte": "", "pages": None,
                                   "confiance": "a_valider"}
        taux = emprise["taux"]
        if emprise["mode"] == "conditionnelle" and code == "UA5a":
            taux = 0.80  # cas « implantation à l'alignement » ; en recul : non réglementée
        calculable = taux is not None and emprise["mode"] not in MODES_NON_CHIFFRES
        if q is None:
            sans_excel.append(code)
        else:
            tx = q.get("emprise_xlsx")
            if (tx is None) != (taux is None) or (tx is not None and abs(tx - taux) > 1e-9):
                ecarts.append({"secteur": code, "xml": taux, "excel": tx})
        fam = famille_de(code, familles)
        sortie[code] = {
            "zone": s["zone"],
            "categorie": s["categorie"],
            "famille": fam,
            "emprise": {
                "taux": taux if calculable else None,
                "calculable": calculable,
                "mode": emprise["mode"],
                "texte": emprise["texte"],
                "pages": emprise["pages"],
                "confiance": emprise["confiance"],
            },
            "piscine": s["piscine"]["texte"] if s["piscine"] else None,
            "habitation": (q or {}).get("habitation"),
            "profil": (q or {}).get("profil"),
            "forme_urbaine": (q or {}).get("forme_urbaine"),
            "hauteur": (q or {}).get("hauteur"),
            "contraintes": (q or {}).get("contraintes"),
            "score_zone": (q or {}).get("score_zone"),
            "priorite_zone": (q or {}).get("priorite_zone"),
            "controle": (q or {}).get("controle"),
            "pages": (q or {}).get("pages_excel") or s["pages_zone"],
        }
    return {
        "source": meta,
        "familles": {k: {kk: vv for kk, vv in v.items() if kk != "secteurs"} for k, v in familles.items()},
        "secteurs": sortie,
        "anomalies": anomalies,
        "controle": {"ecarts_xml_excel": ecarts, "secteurs_sans_excel": sans_excel},
    }


def ecrire_rapport(matrice: dict) -> str:
    c = matrice["controle"]
    lignes = [
        "# Rapport de génération — matrice PLU Sète",
        "",
        f"Source : {matrice['source']['document']} ({matrice['source']['fichier']}).",
        f"Secteurs : {len(matrice['secteurs'])}.",
        "",
        f"- Écarts de taux XML ↔ Excel : **{len(c['ecarts_xml_excel'])}**",
        *[f"  - {e['secteur']} : XML {e['xml']} / Excel {e['excel']}" for e in c["ecarts_xml_excel"]],
        f"- Secteurs sans ligne Excel : {', '.join(c['secteurs_sans_excel']) or 'aucun'}",
        "",
        "## Règles non calculables (réserve affichée « non calculable »)",
        *[f"- {k} ({v['emprise']['mode']}) : {v['emprise']['texte']}"
          for k, v in matrice["secteurs"].items() if not v["emprise"]["calculable"]],
        "",
        "## Anomalies relevées dans l'Excel",
        *[f"- {a['secteurs']} : {a['note']}" for a in matrice["anomalies"]],
        "",
    ]
    return "\n".join(lignes)


def main() -> None:
    matrice = construire()
    (RACINE / "config" / "plu_sete.yaml").write_text(
        "# FICHIER GÉNÉRÉ par pipeline/build_regles.py — ne pas modifier à la main.\n"
        + yaml.safe_dump(matrice, allow_unicode=True, sort_keys=False, width=110),
        encoding="utf-8",
    )
    (RACINE / "web" / "data" / "regles.json").write_text(
        json.dumps(matrice, ensure_ascii=False), encoding="utf-8")
    (RACINE / "config" / "rapport_regles.md").write_text(ecrire_rapport(matrice), encoding="utf-8")
    c = matrice["controle"]
    print(f"{len(matrice['secteurs'])} secteurs ; écarts XML/Excel : {len(c['ecarts_xml_excel'])} ; "
          f"sans Excel : {c['secteurs_sans_excel']}")


if __name__ == "__main__":
    main()
